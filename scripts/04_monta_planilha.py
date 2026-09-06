"""
04. Monta a planilha de entrega.

Entrada: data/processed/
Saída:   outputs/extracao_brasil_2002_2026.xlsx
"""
import os
from comum import PROC, OUT
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

painel = pd.read_csv(os.path.join(PROC,'painel_final.csv'))
res    = pd.read_csv(os.path.join(PROC,'resultados_diferencial.csv'))
cov    = pd.read_csv(os.path.join(PROC,'cobertura.csv'))

painel = painel[['indicador','indicador_nome','iso','pais','ano','valor']].sort_values(
    ['indicador_nome','iso','ano'])

keep = ['indicador','regime','janela','anos_brasil','delta_brasil',
        'delta_America Latina','n_America Latina','dif_America Latina',
        'delta_G20 emergentes','n_G20 emergentes','dif_G20 emergentes',
        'delta_G20 emerg. sem China','n_G20 emerg. sem China','dif_G20 emerg. sem China',
        'delta_Exportadores commodities','n_Exportadores commodities','dif_Exportadores commodities']
res = res[[c for c in keep if c in res.columns]].round(2)

cov['anos_no_periodo'] = cov.obs_reais
cov = cov[['indicador','iso','obs_reais','primeiro','ultimo']].sort_values(['indicador','iso'])

leg = pd.DataFrame({
 'Campo':['delta_brasil','delta_<grupo>','dif_<grupo>','n_<grupo>','anos_brasil','Fonte',
          'Interpolacao','Aviso 1','Aviso 2'],
 'Significado':[
  'Variacao do indicador no Brasil entre o primeiro e o ultimo ano disponivel na janela',
  'Mediana das variacoes dos paises do grupo na mesma janela',
  'Desempenho diferencial = delta_brasil menos delta do grupo. NAO e contrafactual causal.',
  'Quantos paises do grupo tinham dado suficiente. Minimo exigido: 3.',
  'Anos efetivamente usados para o Brasil. Podem diferir da janela nominal do regime.',
  'World Bank API (WDI e PIP), extracao em 2026-09-06, base atualizada em 2026-07-13.',
  'Series irregulares foram interpoladas linearmente APENAS dentro do intervalo observado.',
  'Sinal negativo em pobreza, Gini e homicidios significa MELHORA.',
  'Sinal negativo em renda, PIB e produtividade significa PIORA.']})

os.makedirs(OUT, exist_ok=True)
out = os.path.join(OUT,'extracao_brasil_2002_2026.xlsx')
with pd.ExcelWriter(out, engine='openpyxl') as w:
    leg.to_excel(w, sheet_name='Legenda', index=False)
    res.to_excel(w, sheet_name='Desempenho diferencial', index=False)
    cov.to_excel(w, sheet_name='Cobertura de dados', index=False)
    painel.to_excel(w, sheet_name='Painel completo', index=False)

wb = load_workbook(out)
hdr = Font(name='Arial', bold=True, color='FFFFFF')
fill = PatternFill('solid', fgColor='2F5597')
for ws in wb:
    for c in ws[1]:
        c.font, c.fill, c.alignment = hdr, fill, Alignment(horizontal='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    for col in ws.columns:
        letter = col[0].column_letter
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:400])
        ws.column_dimensions[letter].width = min(max(width + 2, 11), 46)
        for c in col[1:]:
            c.font = Font(name='Arial', size=10)
wb.save(out)
print('ok:', out)
print('abas:', wb.sheetnames)
print('linhas painel:', len(painel), '| linhas resultados:', len(res), '| linhas cobertura:', len(cov))
